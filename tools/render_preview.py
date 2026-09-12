#!/usr/bin/env python3
"""
Render a sheet of a generated .xlsx to a PNG so the design can be eyeballed
without Excel (there is no Excel or LibreOffice in this sandbox).

It is an approximation of Excel's rendering: cached formula values, solid
fills, borders, merged ranges, wrapped text, horizontal/vertical alignment,
number formats and text that spills into empty neighbours.  Emoji are drawn as
coloured tiles because no emoji font is installed - their position and size are
still visible, which is what matters for layout checking.

    /tmp/venv/bin/python tools/render_preview.py products/x.xlsx "🎄 Dashboard" \\
        /tmp/dash.png --rows 1-90
"""

import argparse
import datetime
import re
import sys
import unicodedata

import openpyxl
from openpyxl.utils import get_column_letter, range_boundaries
from PIL import Image, ImageDraw, ImageFont

FONT_DIR = "/usr/share/fonts/truetype/dejavu/"
PX_PER_WIDTH = 7.0
PT_TO_PX = 96.0 / 72.0
DEFAULT_ROW_PX = 20
EMOJI_RE = re.compile(
    "[\U0001F000-\U0001FAFF\u2600-\u27BF\u2B00-\u2BFF\uFE0F\u2190-\u21FF"
    "\u2700-\u27BF\u2764\u2699\u26A1\u2705\u274C\u2753\u2764]")


def is_emoji(ch):
    if ord(ch) < 0x2100:
        return False
    if ord(ch) in (0xFE0F, 0x2022, 0x2014, 0x2013, 0x201C, 0x201D, 0x2019,
                   0x2026, 0x00B7):
        return False
    cat = unicodedata.category(ch)
    return cat in ("So", "Sk", "Cs") or ord(ch) > 0x2500


def font(size, bold=False, italic=False):
    name = "DejaVuSans-Bold.ttf" if bold else "DejaVuSans.ttf"
    return ImageFont.truetype(FONT_DIR + name, max(6, int(round(size))))


def color_of(c, default=None):
    if c is None:
        return default
    rgb = getattr(c, "rgb", None)
    if isinstance(rgb, str) and len(rgb) >= 6:
        rgb = rgb[-6:]
        try:
            return tuple(int(rgb[i:i + 2], 16) for i in (0, 2, 4))
        except ValueError:
            return default
    return default


def fmt_value(value, num_format):
    if value is None:
        return ""
    if isinstance(value, datetime.datetime):
        if value.hour or value.minute:
            return value.strftime("%d %b %Y %H:%M")
        return value.strftime("%d %b %Y")
    if isinstance(value, str):
        return value
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    nf = (num_format or "").lower()
    if isinstance(value, (int, float)):
        if "yy" in nf or "dd" in nf or "mmmm" in nf or "hh" in nf:
            try:
                dt = (datetime.datetime(1899, 12, 30)
                      + datetime.timedelta(days=float(value)))
            except (OverflowError, ValueError):
                return str(value)
            if "yyyy" in nf and "d" not in nf:
                return dt.strftime("%Y")
            if "dddd" in nf:
                return dt.strftime("%A %d %B %Y")
            if "ddd" in nf or "dd mmm" in nf or "d mmm" in nf:
                return dt.strftime("%d %b %Y")
            if "mm" in nf or "yy" in nf:
                return dt.strftime("%d/%m/%Y")
            return dt.strftime("%d %b")
        if "%" in nf:
            dp = 2 if "0.00%" in nf else (1 if "0.0%" in nf else 0)
            return ("%%.%df%%%%" % dp) % (value * 100)
        if "#,##0.00" in nf:
            return "{:,.2f}".format(value)
        if "#,##0" in nf or nf in ("0", "#,##0.0"):
            return "{:,.0f}".format(value) if "#,##0.0" not in nf \
                else "{:,.1f}".format(value)
        if value == int(value):
            return str(int(value))
        return "%.2f" % value
    return str(value)


def render(path, sheet_title, out, rows=None, cols=None, zoom=1.0):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet_title]

    # --- geometry -------------------------------------------------------
    widths = {}
    for dim in ws.column_dimensions.values():
        if not dim.width:
            continue
        for idx in range(dim.min or 1, (dim.max or dim.min or 1) + 1):
            widths[idx] = dim.width
    heights = {i: d.height for i, d in ws.row_dimensions.items() if d.height}

    r1, r2 = (rows or (1, ws.max_row))
    c1, c2 = (cols or (1, ws.max_column))

    def col_px(i):
        return int(round(widths.get(i, 8.43) * PX_PER_WIDTH)) + 5

    def row_px(i):
        return int(round(heights.get(i, 15.0) * PT_TO_PX)) or DEFAULT_ROW_PX

    xs = {}
    x = 0
    for i in range(c1, c2 + 1):
        xs[i] = x
        x += col_px(i)
    ys = {}
    y = 0
    for i in range(r1, r2 + 1):
        ys[i] = y
        y += row_px(i)
    W, H = int(x * zoom), int(y * zoom)
    img = Image.new("RGB", (max(W, 1), max(H, 1)), (255, 255, 255))
    d = ImageDraw.Draw(img)

    merges = {}
    swallowed = set()
    for mr in ws.merged_cells.ranges:
        mc1, mr1, mc2, mr2 = range_boundaries(str(mr))
        merges[(mr1, mc1)] = (mr1, mc1, mr2, mc2)
        for rr in range(mr1, mr2 + 1):
            for cc in range(mc1, mc2 + 1):
                if (rr, cc) != (mr1, mc1):
                    swallowed.add((rr, cc))

    def cell_box(rr, cc):
        if (rr, cc) in merges:
            a, b, c_, e = merges[(rr, cc)]
            x1, y1 = xs.get(b, x), ys.get(a, y)
            x2 = xs.get(e, x) + col_px(e)
            y2 = ys.get(c_, y) + row_px(c_)
        else:
            x1, y1 = xs.get(cc, x), ys.get(rr, y)
            x2, y2 = x1 + col_px(cc), y1 + row_px(rr)
        return x1, y1, x2, y2

    def occupied(rr, cc):
        c = ws.cell(row=rr, column=cc)
        return c.value not in (None, "")

    # --- paint ----------------------------------------------------------
    for rr in range(r1, r2 + 1):
        for cc in range(c1, c2 + 1):
            if (rr, cc) in swallowed:
                continue
            cell = ws.cell(row=rr, column=cc)
            x1, y1, x2, y2 = cell_box(rr, cc)
            if x2 <= 0 or y2 <= 0 or x1 >= x or y1 >= y:
                continue
            fill = color_of(cell.fill.start_color) if (
                cell.fill and cell.fill.patternType == "solid") else None
            if fill:
                d.rectangle([x1, y1, x2 - 1, y2 - 1], fill=fill)
            b = cell.border
            for side, edges in ((b.left, (x1, y1, x1, y2)),
                                (b.right, (x2 - 1, y1, x2 - 1, y2)),
                                (b.top, (x1, y1, x2, y1)),
                                (b.bottom, (x1, y2 - 1, x2, y2 - 1))):
                if side is not None and side.style:
                    col = color_of(side.color, (200, 200, 200))
                    thick = 2 if side.style in ("medium", "thick", "double") \
                        else 1
                    d.line(list(edges), fill=col, width=thick)

    # --- text -----------------------------------------------------------
    for rr in range(r1, r2 + 1):
        for cc in range(c1, c2 + 1):
            if (rr, cc) in swallowed:
                continue
            cell = ws.cell(row=rr, column=cc)
            raw = fmt_value(cell.value, cell.number_format)
            if not raw:
                continue
            x1, y1, x2, y2 = cell_box(rr, cc)
            if x2 <= 0 or y2 <= 0 or x1 >= x or y1 >= y:
                continue
            al = cell.alignment
            fnt = font((cell.font.size or 11) * 1.02,
                       bool(cell.font.bold), bool(cell.font.italic))
            col = color_of(cell.font.color, (0, 0, 0))
            indent = int((al.indent or 0) * 8)
            pad = 4 + indent

            # spill into empty neighbours when the text is not wrapped
            box_w = x2 - x1
            if not al.wrap_text:
                cc2 = cc + 1
                while (cc2 <= c2 and not occupied(rr, cc2)
                       and (rr, cc2) not in swallowed
                       and d.textlength(raw, font=fnt) > box_w - 2 * pad):
                    box_w += col_px(cc2)
                    cc2 += 1

            if al.wrap_text:
                words, lines, cur = raw.replace("\n", " \n ").split(), [], ""
                for w in words:
                    trial = (cur + " " + w).strip()
                    if d.textlength(trial, font=fnt) <= box_w - 2 * pad \
                            or not cur:
                        cur = trial
                    else:
                        lines.append(cur)
                        cur = w
                    if w == "\n":
                        lines.append(cur)
                        cur = ""
                if cur:
                    lines.append(cur)
            else:
                lines = raw.split("\n")

            lh = int((cell.font.size or 11) * 1.42)
            total_h = lh * len(lines)
            va = al.vertical or "bottom"
            if va == "center":
                ty = y1 + max(0, (y2 - y1 - total_h) // 2)
            elif va == "top":
                ty = y1 + 3
            else:
                ty = y1 + max(0, (y2 - y1 - total_h) // 2)

            for line in lines:
                drawn, cx = [], x1 + pad
                for ch in line:
                    if is_emoji(ch):
                        drawn.append(("tile", ch))
                    else:
                        drawn.append(("text", ch))
                # measure
                widths_px, tiles = [], []
                for kind, ch in drawn:
                    if kind == "tile":
                        wpx = int(lh * 0.95)
                        tiles.append(wpx)
                    else:
                        wpx = d.textlength(ch, font=fnt)
                    widths_px.append(wpx)
                line_w = sum(widths_px)
                ha = al.horizontal or ("right" if isinstance(
                    cell.value, (int, float)) and not cell.font.bold
                    else "left")
                if ha == "center":
                    cx = x1 + max(pad, (box_w - line_w) // 2)
                elif ha == "right":
                    cx = x1 + max(pad, box_w - line_w - pad)
                for (kind, ch), wpx in zip(drawn, widths_px):
                    if kind == "tile":
                        hue = (ord(ch) * 47) % 360
                        import colorsys
                        rgb = tuple(int(255 * v) for v in
                                    colorsys.hsv_to_rgb(hue / 360.0, 0.55, 0.85))
                        d.rounded_rectangle([cx, ty + 1, cx + wpx - 2,
                                             ty + lh - 2], radius=3, fill=rgb)
                    else:
                        d.text((cx, ty), ch, font=fnt, fill=col)
                    cx += wpx
                ty += lh
    # --- inserted pictures ---------------------------------------------
    from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, \
        TwoCellAnchor, AbsoluteAnchor
    for pic in getattr(ws, "_images", []):
        anc = pic.anchor
        frm = getattr(anc, "_from", None)
        if frm is None:
            continue
        col, row = frm.col + 1, frm.row + 1
        if not (c1 <= col <= c2 and r1 <= row <= r2):
            continue
        x1 = xs.get(col, 0) + int((frm.colOff or 0) / 9525.0)
        y1 = ys.get(row, 0) + int((frm.rowOff or 0) / 9525.0)
        to = getattr(anc, "to", None)
        if to is not None:
            x2 = xs.get(to.col + 1, x) + int((to.colOff or 0) / 9525.0)
            y2 = ys.get(to.row + 1, y) + int((to.rowOff or 0) / 9525.0)
            wpx, hpx = x2 - x1, y2 - y1
        else:
            try:
                ext = anc.ext
                wpx, hpx = int(ext.cx / 9525.0), int(ext.cy / 9525.0)
            except AttributeError:
                wpx, hpx = int(pic.width), int(pic.height)
        try:
            import io as _io
            sub = Image.open(_io.BytesIO(pic._data())).convert("RGBA")
            sub = sub.resize((max(wpx, 1), max(hpx, 1)), Image.LANCZOS)
            img.paste(sub, (x1, y1), sub)
        except Exception as exc:
            print("  (image skipped: %s)" % exc)

    if zoom != 1.0:
        img = img.resize((W, H), Image.LANCZOS)
    img.save(out)
    return out, img.size


def main(argv):
    p = argparse.ArgumentParser()
    p.add_argument("workbook")
    p.add_argument("sheet")
    p.add_argument("out")
    p.add_argument("--rows", default=None, help="e.g. 1-90")
    p.add_argument("--cols", default=None, help="e.g. A-M")
    p.add_argument("--zoom", type=float, default=1.0)
    a = p.parse_args(argv)
    rows = tuple(int(x) for x in a.rows.split("-")) if a.rows else None
    cols = None
    if a.cols:
        lo, hi = a.cols.split("-")
        cols = (openpyxl.utils.column_index_from_string(lo),
                openpyxl.utils.column_index_from_string(hi))
    out, size = render(a.workbook, a.sheet, a.out, rows, cols, a.zoom)
    print("%s -> %s (%dx%d)" % (a.sheet, out, size[0], size[1]))
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
