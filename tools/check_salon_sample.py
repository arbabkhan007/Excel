#!/usr/bin/env python3
"""Does the demo salon actually hang together?

The validator checks the formulas; this checks the *story* they are
told. It reads only the literal cells a user would have typed - never
a formula column - rebuilds the arithmetic in Python, and complains if
the example data would make the workbook look wrong on the first open:
double bookings, visits with nothing rung up, payments that do not
match the till, commission paid that was never earned, stock that goes
negative, IDs that point at nothing.

    python tools/check_salon_sample.py output/salon.xlsx
"""

from __future__ import annotations

import datetime as _dt
import os
import re
import sys
from collections import Counter, defaultdict

try:
    import openpyxl
except ImportError:                                   # pragma: no cover
    print("openpyxl is required: pip install openpyxl")
    raise SystemExit(2)

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from salon_business_spreadsheet.build import REGISTER_MODULES  # noqa: E402
from salon_business_spreadsheet.config import (COMMISSION_TIERS,  # noqa: E402
                                               Config, rid)

MONEY = 0.015           # a cent and a half of rounding slack


def _specs(cfg: Config):
    return {module.NAME: module.spec(cfg) for module in REGISTER_MODULES}


def _rows(ws, spec, key_column: str):
    """Every filled data row as {column key: literal value}."""
    first = spec.first_data_row                      # 1 based
    last = spec.last_data_row
    key_index = spec.index(key_column) + 1
    out = []
    for excel_row in range(first, last + 1):
        anchor = ws.cell(row=excel_row, column=key_index).value
        if anchor in (None, ""):
            continue
        record = {"_row": excel_row - first + 1}
        for index, col in enumerate(spec.columns):
            value = ws.cell(row=excel_row, column=index + 1).value
            if isinstance(value, str) and value.startswith("="):
                continue                              # never trust a formula
            record[col.key] = value
        out.append(record)
    return out


def _month(value) -> str:
    if isinstance(value, _dt.datetime):
        value = value.date()
    return f"{value.year:04d}-{value.month:02d}" if value else ""


def _day(value):
    return value.date() if isinstance(value, _dt.datetime) else value


def _minutes(value) -> int:
    if isinstance(value, _dt.datetime):
        return value.hour * 60 + value.minute
    if isinstance(value, _dt.time):
        return value.hour * 60 + value.minute
    if isinstance(value, (int, float)):
        return int(round(float(value) * 24 * 60))
    return 0


def check(path: str) -> int:                          # noqa: C901
    cfg = Config()
    specs = _specs(cfg)
    wb = openpyxl.load_workbook(path, data_only=False)
    problems: list = []
    notes: list = []

    data = {}
    anchors = {
        "Services": "name", "Clients": "name", "Staff": "name",
        "Appointments": "date", "Sales": "date", "Payments": "date",
        "Commissions": "month", "Payroll": "period_end",
        "Expenses": "date", "Inventory": "name", "Purchases": "date",
        "Marketing": "name", "Loyalty": "client",
    }
    for name, anchor in anchors.items():
        if name not in wb.sheetnames:
            problems.append(f"missing sheet {name}")
            continue
        data[name] = _rows(wb[name], specs[name], anchor)
    if problems:
        for line in problems:
            print("    - " + line)
        return 1

    counts = ", ".join(f"{name} {len(rows)}" for name, rows in data.items())
    notes.append("demo rows: " + counts)
    if not data["Appointments"]:
        print(f"--- {path}\n  . no example data in this build (blank "
              f"version) - nothing to check")
        return 0

    # ---------------------------------------------------- the ID spine --
    service_by_id = {rid("Service", row["_row"]): row
                     for row in data["Services"]}
    client_by_id = {rid("Client", row["_row"]): row
                    for row in data["Clients"]}
    staff_by_id = {rid("Staff", row["_row"]): row for row in data["Staff"]}
    product_by_id = {rid("Product", row["_row"]): row
                     for row in data["Inventory"]}
    appt_by_id = {rid("Appointment", row["_row"]): row
                  for row in data["Appointments"]}
    sale_by_id = {rid("Sale", row["_row"]): row for row in data["Sales"]}

    def known(table, value, where, kind):
        if value in (None, ""):
            return False
        if value not in table:
            problems.append(f"{where}: {kind} '{value}' does not exist")
            return False
        return True

    for row in data["Appointments"]:
        where = f"Appointments row {row['_row']}"
        known(client_by_id, row.get("client"), where, "client")
        known(service_by_id, row.get("service"), where, "service")
        known(staff_by_id, row.get("staff"), where, "stylist")
    for row in data["Sales"]:
        where = f"Sales row {row['_row']}"
        if row.get("appointment"):
            known(appt_by_id, row["appointment"], where, "appointment")
        if row.get("client"):
            known(client_by_id, row["client"], where, "client")
        if row.get("staff"):
            known(staff_by_id, row["staff"], where, "stylist")
        if row.get("kind") == "Service":
            known(service_by_id, row.get("item"), where, "service")
        if row.get("kind") == "Retail":
            known(product_by_id, row.get("item"), where, "product")
    for row in data["Payments"]:
        where = f"Payments row {row['_row']}"
        if row.get("sale"):
            known(sale_by_id, row["sale"], where, "sale")
        if row.get("appointment"):
            known(appt_by_id, row["appointment"], where, "appointment")

    # ------------------------------------------------- double bookings --
    diary = defaultdict(list)
    for row in data["Appointments"]:
        if row.get("status") == "Cancelled":
            continue
        service = service_by_id.get(row.get("service"), {})
        minutes = row.get("minutes") or service.get("duration") or 0
        buffer_ = service.get("buffer") or 0
        start = _minutes(row.get("start"))
        finish = start + int(minutes) + int(buffer_)
        diary[(row.get("staff"), _day(row.get("date")))].append(
            (start, finish, row["_row"]))
    clashes = 0
    for (who, day), slots in diary.items():
        slots.sort()
        for first, second in zip(slots, slots[1:]):
            if second[0] < first[1]:
                clashes += 1
                if clashes < 6:
                    problems.append(
                        f"double booking: {who} on {day}, rows "
                        f"{first[2]} and {second[2]}")
    notes.append(f"diary: {len(data['Appointments'])} bookings, "
                 f"{clashes} clashes")

    # --------------------------------------------- every visit rung up --
    sales_by_appt = defaultdict(list)
    for row in data["Sales"]:
        if row.get("appointment"):
            sales_by_appt[row["appointment"]].append(row)
    missing_sale = 0
    for ref, row in appt_by_id.items():
        if row.get("status") != "Completed":
            continue
        if row.get("price") == 0:
            continue                                  # chair renter
        if not sales_by_appt.get(ref):
            missing_sale += 1
    if missing_sale:
        problems.append(f"{missing_sale} completed visits have no sale "
                        f"logged against them")

    # ----------------------------------------------- money against till --
    def sale_total(row):
        kind = row.get("kind")
        qty = row.get("qty") or 1
        if row.get("unit_price") not in (None, ""):
            price = float(row["unit_price"]) * qty
        elif kind == "Service":
            price = float(service_by_id.get(row.get("item"), {})
                          .get("price") or 0)
        elif kind == "Retail":
            price = float(product_by_id.get(row.get("item"), {})
                          .get("retail_price") or 0) * qty
        else:
            price = 0.0
        discount = round(price * float(row.get("discount_pct") or 0), 2) + \
            float(row.get("discount_amt") or 0)
        net = round(price - discount - float(row.get("refund") or 0), 2)
        taxed = ((kind == "Retail" and cfg.tax_on_retail == "Yes") or
                 (kind != "Retail" and cfg.tax_on_services == "Yes"))
        tax = round(net * cfg.tax_rate, 2) if taxed else 0.0
        return round(net + tax, 2), net

    paid_by_sale = Counter()
    deposits_by_appt = Counter()
    tips = 0.0
    for row in data["Payments"]:
        amount = float(row.get("amount") or 0)
        kind = row.get("kind")
        if kind == "Tip":
            tips += amount
            continue
        if kind == "Refund":
            amount = -amount
        if row.get("sale"):
            paid_by_sale[row["sale"]] += amount
        elif row.get("appointment"):
            deposits_by_appt[row["appointment"]] += amount
    first_of_appt = {}
    for ref, row in sale_by_id.items():
        appt = row.get("appointment")
        if appt and appt not in first_of_appt:
            first_of_appt[appt] = ref
    unpaid = 0
    overpaid = 0
    net_sales = 0.0
    gross_by_month = Counter()
    for ref, row in sale_by_id.items():
        total, net = sale_total(row)
        net_sales += net
        gross_by_month[_month(row.get("date"))] += net
        paid = paid_by_sale.get(ref, 0.0)
        if first_of_appt.get(row.get("appointment")) == ref:
            paid += deposits_by_appt.get(row["appointment"], 0.0)
        balance = round(total - paid, 2)
        if balance > MONEY:
            unpaid += 1
        elif balance < -MONEY:
            overpaid += 1
            if overpaid < 6:
                problems.append(
                    f"Sales row {row['_row']}: paid {paid:,.2f} against "
                    f"a bill of {total:,.2f} - money applied twice?")
    notes.append(f"till: {len(sale_by_id)} sales, net {net_sales:,.0f}, "
                 f"{unpaid} left unpaid on purpose, {overpaid} overpaid")
    if unpaid > len(sale_by_id) * 0.1:
        problems.append(f"{unpaid} unpaid sales is too many for a demo "
                        f"({len(sale_by_id)} sales)")

    # ------------------------------------------------------- commission --
    by_staff_month = defaultdict(lambda: dict(service=0.0, retail=0.0,
                                              cost=0.0, count=0))
    for ref, row in sale_by_id.items():
        who = row.get("staff")
        if not who:
            continue
        bucket = by_staff_month[(who, _month(row.get("date")))]
        _total, net = sale_total(row)
        kind = row.get("kind")
        if kind == "Service":
            service = service_by_id.get(row.get("item"), {})
            bucket["service"] += net
            bucket["cost"] += float(service.get("product_cost") or 0)
            bucket["count"] += 1
        elif kind == "Retail":
            bucket["retail"] += net

    def tier_rate(base: float) -> float:
        rate = COMMISSION_TIERS[0][1]
        for floor, value in COMMISSION_TIERS:
            if base >= floor:
                rate = value
        return rate

    def commission_for(who: str, month: str) -> float:
        person = staff_by_id.get(who, {})
        bucket = by_staff_month.get((who, month))
        if not bucket:
            return 0.0
        method = person.get("method") or "None"
        service = round(bucket["service"], 2)
        retail = round(bucket["retail"], 2)
        cost = round(bucket["cost"], 2)
        rate = float(person.get("rate_service") or cfg.service_commission)
        retail_rate = float(person.get("rate_retail") or
                            cfg.retail_commission)
        if method == "Percent of service":
            return round(service * rate, 2)
        if method == "Percent of service and retail":
            return round(round(service * rate, 2) +
                         round(retail * retail_rate, 2), 2)
        if method == "Percent after product cost":
            return round(max(0.0, service - cost) * rate, 2)
        if method == "Tiered on monthly service revenue":
            return round(service * tier_rate(service), 2)
        if method == "Flat amount per service":
            return round(bucket["count"] *
                         float(person.get("flat_rate") or 0), 2)
        return 0.0

    seen = Counter()
    for row in data["Commissions"]:
        key = (row.get("staff"), row.get("month"))
        seen[key] += 1
        if seen[key] > 1:
            problems.append(f"Commissions row {row['_row']}: a second row "
                            f"for {key[0]} in {key[1]}")
        if row.get("staff") not in staff_by_id:
            problems.append(f"Commissions row {row['_row']}: unknown "
                            f"stylist {row.get('staff')}")
        if not re.match(r"^\d{4}-\d{2}$", str(row.get("month") or "")):
            problems.append(f"Commissions row {row['_row']}: month "
                            f"'{row.get('month')}' is not yyyy-mm")
    methods = {staff_by_id[k[0]].get("method")
               for k in seen if k[0] in staff_by_id}
    notes.append(f"commissions: {len(data['Commissions'])} rows covering "
                 f"{len(methods)} different pay methods")

    released = Counter()
    for row in data["Payroll"]:
        month = _month(row.get("period_end"))
        if row.get("commission"):
            released[(row.get("staff"), month)] += float(row["commission"])
        if row.get("staff") not in staff_by_id:
            problems.append(f"Payroll row {row['_row']}: unknown stylist")
        if not row.get("hours") and not row.get("commission"):
            problems.append(f"Payroll row {row['_row']}: neither hours "
                            f"nor commission")
    this_month = f"{cfg.today.year:04d}-{cfg.today.month:02d}"
    for (who, month), amount in released.items():
        if month == this_month:
            continue
        earned = commission_for(who, month)
        adjust = sum(float(row.get("adjust") or 0) + float(
            row.get("bonus") or 0)
            for row in data["Commissions"]
            if row.get("staff") == who and row.get("month") == month)
        if amount - (earned + adjust) > 1.0:
            problems.append(
                f"payroll releases {amount:,.2f} of commission to {who} "
                f"in {month} but only {earned + adjust:,.2f} was earned")
    notes.append(f"payroll: {len(data['Payroll'])} runs, "
                 f"{len(released)} stylist-months with commission")

    # ------------------------------------------------------------ stock --
    received = Counter()
    for row in data["Purchases"]:
        status = row.get("status")
        qty = float(row.get("qty") or 0)
        got = row.get("qty_received")
        if status == "Received":
            received[row.get("product")] += float(got) if got else qty
        elif status == "Part received":
            received[row.get("product")] += float(got or 0)
        if row.get("product") not in product_by_id:
            problems.append(f"Purchases row {row['_row']}: unknown product")
    sold = Counter()
    for row in data["Sales"]:
        if row.get("kind") == "Retail":
            sold[row.get("item")] += float(row.get("qty") or 1)
    negative = 0
    for ref, row in product_by_id.items():
        stock = (float(row.get("opening") or 0) + received[ref] -
                 float(row.get("used") or 0) - sold[ref] -
                 float(row.get("waste") or 0) +
                 float(row.get("adjust") or 0))
        if stock < 0:
            negative += 1
            if negative < 6:
                problems.append(
                    f"Inventory row {row['_row']} ({row.get('name')}): "
                    f"stock works out at {stock:,.1f}")
    notes.append(f"stock: {len(product_by_id)} products, {negative} would "
                 f"go negative")

    # -------------------------------------------------------- marketing --
    codes = [row.get("code") for row in data["Marketing"] if row.get("code")]
    duplicates = [code for code, n in Counter(codes).items() if n > 1]
    if duplicates:
        problems.append(f"marketing codes used twice: {duplicates}")
    used_codes = {row.get("promo") for row in data["Sales"]
                  if row.get("promo")}
    unknown = used_codes - set(codes)
    if unknown:
        problems.append(f"sales use promo codes that are not on the "
                        f"marketing tab: {sorted(unknown)}")
    notes.append(f"marketing: {len(codes)} codes, {len(used_codes)} of "
                 f"them actually used on sales")

    members = [row.get("client") for row in data["Loyalty"]]
    repeats = [c for c, n in Counter(members).items() if n > 1]
    if repeats:
        problems.append(f"loyalty: {len(repeats)} clients listed twice")

    # ------------------------------------------------------- the story --
    months = sorted(gross_by_month)
    recent = [m for m in months if m < this_month][-3:]
    if recent:
        average = sum(gross_by_month[m] for m in recent) / len(recent)
        notes.append("net sales by month: " + ", ".join(
            f"{m} {gross_by_month[m]:,.0f}" for m in months[-4:]))
        target = cfg.revenue_target_month
        if average < target * 0.6 or average > target * 1.6:
            problems.append(
                f"the demo averages {average:,.0f} a month against a "
                f"target of {target:,.0f} - the dashboard will look odd")
    completed = sum(1 for row in data["Appointments"]
                    if row.get("status") == "Completed")
    final = sum(1 for row in data["Appointments"]
                if row.get("status") in ("Completed", "Cancelled",
                                         "No show"))
    if final:
        notes.append(f"completion rate {completed / final:.0%}, "
                     f"average ticket "
                     f"{net_sales / max(1, completed):,.2f}, "
                     f"tips {tips:,.0f}")

    print(f"--- {path}")
    for note in notes:
        print(f"  . {note}")
    if problems:
        print(f"  ! {len(problems)} problem(s):")
        for line in problems[:30]:
            print("    - " + line)
        return 1
    print("  OK - the example salon adds up")
    return 0


if __name__ == "__main__":
    target = (sys.argv[1] if len(sys.argv) > 1
              else "output/Salon Business Management Spreadsheet.xlsx")
    raise SystemExit(check(target))
